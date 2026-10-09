import csv
import uuid
import zipfile
from types import SimpleNamespace

from openpyxl import load_workbook

from app.config import settings
from app.services.export import engine, storage


class Client:
    def __init__(self, rows):
        self.rows = rows
        self.devices = object()
        self.vms = object()
        self.vdcs = object()
        self.sites = object()
        self.device_types = object()
        self.nb = SimpleNamespace(
            dcim=SimpleNamespace(devices=self.devices, virtual_device_contexts=self.vdcs,
                                 sites=self.sites, device_types=self.device_types),
            virtualization=SimpleNamespace(virtual_machines=self.vms),
        )

    def paginated(self, endpoint, on_page=None, **filters):
        values = self.rows.get(endpoint, [])
        if on_page:
            on_page({"count": len(values), "results": values})
        yield from values


def spec(format="csv", types=None, delimiter=","):
    return engine.JobSpec(str(uuid.uuid4()), 1, types or ["device"], {
        "device": {"optional": [], "custom_fields": [{"name": "formula", "type": "text"}, {"name": "flag", "type": "boolean"}]},
        "virtualmachine": {"optional": [], "custom_fields": []},
    }, format, delimiter)


def test_safe_cell_blocks_formulas_but_not_numbers():
    assert engine.safe_cell("=cmd") == "'=cmd"
    assert engine.safe_cell("-12.5") == "-12.5"
    assert engine.safe_cell("+12") == "+12"


def test_csv_bom_delimiter_and_custom_values(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "export_dir", str(tmp_path))
    storage.ensure_dirs()
    client = Client({})
    client.rows[client.devices] = [{"id": 1, "name": "d1", "primary_ip": {"address": "1.2.3.4/32"}, "serial": "s", "custom_fields": {"formula": "=bad", "flag": True}}]
    result = engine.run_export(spec(delimiter=";"), client, lambda *args: None, lambda: False)
    raw = (tmp_path / "files" / result.file_name).read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(raw.decode("utf-8-sig").splitlines(), delimiter=";"))
    assert rows[0] == ["name", "primary_ip", "serial", "cf_formula", "cf_flag"]
    assert rows[1][-2:] == ["'=bad", "true"]


def test_xlsx_and_multitype_zip(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "export_dir", str(tmp_path))
    storage.ensure_dirs()
    client = Client({})
    client.rows[client.devices] = [{"id": 1, "name": "d", "custom_fields": {}}]
    client.rows[client.vms] = [{"id": 2, "name": "vm", "custom_fields": {}}]
    xlsx = engine.run_export(spec("xlsx", ["device", "virtualmachine"]), client, lambda *args: None, lambda: False)
    workbook = load_workbook(tmp_path / "files" / xlsx.file_name, read_only=True)
    assert workbook.sheetnames == ["Devices", "VirtualMachines"]
    zipped = engine.run_export(spec("csv", ["device", "virtualmachine"]), client, lambda *args: None, lambda: False)
    with zipfile.ZipFile(tmp_path / "files" / zipped.file_name) as archive:
        assert set(archive.namelist()) == {"devices.csv", "virtual-machines.csv"}


def test_cancel_cleans_tmp_files(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "export_dir", str(tmp_path))
    storage.ensure_dirs()
    client = Client({})
    client.rows[client.devices] = [{"id": 1, "name": "d", "custom_fields": {}}]
    try:
        engine.run_export(spec(), client, lambda *args: None, lambda: True)
    except engine.ExportCancelled:
        pass
    assert list((tmp_path / "tmp").iterdir()) == []
